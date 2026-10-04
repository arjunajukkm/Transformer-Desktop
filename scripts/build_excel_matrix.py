# -*- coding: utf-8 -*-
"""
build_excel_matrix.py
─────────────────────
Generates the Workforce Trends Metrics Logic Excel workbook matching
the exact 5-column structure requested by the user:
  Row 1: Matrix Name | Trajactory and Benchmark | Business unit benchmark ranking | Business Unit Trend Heatmap | What changes
  Row 2: Name        | Logic                    | Logic                           | Logic                       | Logic
"""

from pathlib import Path
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

def generate_workbook(out_path: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Trends Metrics Logic"
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions
    c_navy_dark = "0F172A"
    c_navy_light = "1E293B"
    c_zebra_even = "F8FAFC"
    c_zebra_odd = "FFFFFF"
    c_border = "CBD5E1"
    c_blue_title = "0284C7"

    font_r1 = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_r2 = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_metric = Font(name="Segoe UI", size=10, bold=True, color=c_blue_title)
    font_text = Font(name="Segoe UI", size=9.5, color="1E293B")

    fill_r1 = PatternFill(start_color=c_navy_dark, end_color=c_navy_dark, fill_type="solid")
    fill_r2 = PatternFill(start_color=c_navy_light, end_color=c_navy_light, fill_type="solid")

    border_thin = Border(
        left=Side(style="thin", color=c_border),
        right=Side(style="thin", color=c_border),
        top=Side(style="thin", color=c_border),
        bottom=Side(style="thin", color=c_border),
    )

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="top", wrap_text=True)

    headers_row1 = [
        "Matrix Name",
        "Trajactory and Benchmark",
        "Business unit benchmark ranking",
        "Business Unit Trend Heatmap",
        "What changes",
    ]
    headers_row2 = [
        "Name",
        "Logic",
        "Logic",
        "Logic",
        "Logic",
    ]

    ws.append(headers_row1)
    ws.append(headers_row2)

    for c in range(1, 6):
        cell1 = ws.cell(row=1, column=c)
        cell2 = ws.cell(row=2, column=c)

        cell1.font = font_r1
        cell1.fill = fill_r1
        cell1.alignment = align_center
        cell1.border = border_thin

        cell2.font = font_r2
        cell2.fill = fill_r2
        cell2.alignment = align_center
        cell2.border = border_thin

    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 22

    # 18 Metrics Data
    rows_data = [
        (
            "Leave Application Compliance\n(Category: Compliance | ID: leave_application_compliance | Higher is Better)",
            "- Monthly Aggregation: Evaluates all leave requests whose start date falls within the calendar month.\n"
            "- Numerator: Count of leave requests classified as compliant (benchmark_compliant before Oct 1, 2026; policy_compliant on/after Oct 1, 2026).\n"
            "- Denominator: Total evaluable leave requests (total requests minus records flagged with DATA_QUALITY_UNCERTAIN).\n"
            "- Rate Formula: (Compliant Requests / Evaluable Requests) * 100%.\n"
            "- Benchmark Line: Organization-wide average, historical baseline (median pre-policy rate), or peer BU average.\n"
            "- Trajectory Slope: Evaluates trailing 3-month linear slope (beta) vs materiality threshold (1.0 pp). Improving if beta >= +1.0 pp, Deteriorating if beta <= -1.0 pp, else Stable.\n"
            "- Regression Flag: Triggers when prior improvement is reversed by recent deterioration (>= 80% reversal).\n"
            "- Confidence: High (>= 50 requests), Moderate (10-49), Low (< 10).",

            "- Scope & Grouping: Groups evaluated leave requests by Business Unit for the latest reporting month.\n"
            "- Current BU Rate: Compliant Requests / Evaluable Requests for the selected BU.\n"
            "- 3-Month Rolling Average: Arithmetic mean of the BU's monthly compliance rates over the last 3 available months.\n"
            "- Historical Baseline: Median compliance rate of the BU during historical pre-policy periods.\n"
            "- Gap vs Benchmark: (BU Current Rate) - (Org Benchmark Rate) in percentage points (pp). Positive gap indicates outperformance (favorable green badge).\n"
            "- Sorting & Ranking: BUs are ranked from highest compliance to lowest compliance.\n"
            "- Total Row: Organization-wide weighted total compliance rate displayed at table footer.",

            "- Matrix Structure: Rows = Business Units, Columns = Chronological monthly periods (e.g., May 2026, Jun 2026, Jul 2026).\n"
            "- Cell Value: Monthly compliance rate for each BU (e.g., 88.5%).\n"
            "- Delta to Org: Displays difference relative to the organization-wide average for that specific month (e.g., +4.2 pp).\n"
            "- Conditional Shading:\n"
            "   * Green (Good): Compliance >= 85.0% or >= +5.0 pp above company average.\n"
            "   * Amber (Warning): Compliance 70.0%-84.9% or within +/- 5.0 pp of company average.\n"
            "   * Red (Critical): Compliance < 70.0% or > 5.0 pp below company average.\n"
            "   * Gray (N/A): Evaluated units < 5 (insufficient sample size).",

            "- Month-over-Month (MoM) Delta: Total Company Compliance in latest month minus previous month (e.g., +3.2 pp).\n"
            "- Headline: Deterministic insight stating magnitude, direction (increased/decreased), and baseline movement.\n"
            "- Driver Decomposition Formula:\n"
            "   * For each BU: BU Delta = (BU Current Rate) - (BU Previous Rate).\n"
            "   * Volume Weight = (BU Evaluable Requests in Current Month) / (Total Org Evaluable Requests).\n"
            "   * Net Contribution = BU Delta * Volume Weight.\n"
            "- Primary Drivers: Top 4 BUs by absolute movement are isolated, tagged with Favorable or Adverse impact. Remaining BUs aggregated into 'Other Business Units'."
        ),

        (
            "PL Application Compliance\n(Category: Compliance | ID: pl_application_compliance | Higher is Better)",
            "- Scope: Filters to leave requests where is_pl == True (Privilege Leave / Earned Leave).\n"
            "- Policy Rule (Post Oct 1, 2026): Advance notice requirement scales with leave duration:\n"
            "   * Leave duration <= 3 units: Requires >= 2 days advance application notice.\n"
            "   * Leave duration > 3 and <= 7 units: Requires >= 15 days advance application notice.\n"
            "   * Leave duration > 7 units: Requires >= 30 days advance application notice.\n"
            "- Pre-Policy Benchmark: Standard >= 2 days advance notice.\n"
            "- Formula: (Compliant PL Requests / Evaluable PL Requests) * 100%.\n"
            "- Benchmark: Org-wide PL compliance rate or pre-policy historical baseline.",

            "- Aggregation: Computes current month PL compliance per Business Unit.\n"
            "- 3-Month Average: Rolling 3-month mean PL compliance.\n"
            "- Historical: Median pre-policy PL compliance for the BU.\n"
            "- Gap: BU PL Compliance - Org PL Compliance (pp).\n"
            "- Ranking: Sorted descending by current PL compliance rate.",

            "- Cell Value: BU monthly PL compliance percentage.\n"
            "- Color Coding:\n"
            "   * Green: PL Compliance >= 90.0%.\n"
            "   * Amber: PL Compliance 75.0%-89.9%.\n"
            "   * Red: PL Compliance < 75.0% (indicates short-notice long-duration leaves).",

            "- MoM Net Movement: Latest PL compliance % - Prior month PL compliance %.\n"
            "- Driver Attribution: Identifies which BUs submitted planned vacation leaves with sufficient advance notice vs emergency short-notice filings.\n"
            "- Contribution: Weighted by BU's share of total company PL applications."
        ),

        (
            "Non-PL Application Compliance\n(Category: Compliance | ID: non_pl_application_compliance | Higher is Better)",
            "- Scope: Filters to leave requests where is_pl == False (Casual, Sick, Bereavement, Paternity, etc.).\n"
            "- Policy Rule: Post-facto applications must be submitted within 3 calendar days of availing leave (lead_time_days >= -3).\n"
            "- Formula: (Compliant Non-PL Requests / Evaluable Non-PL Requests) * 100%.\n"
            "- Excludes uncertain data quality records.\n"
            "- Trajectory: 3-month slope with 1.0 pp materiality threshold.",

            "- Evaluates adherence to the 3-day post-facto submission window across Business Units.\n"
            "- Gap: BU Non-PL Rate - Org Non-PL Rate (pp).\n"
            "- Ranking: Sorted descending by compliance rate.",

            "- Heatmap Matrix: Monthly Non-PL compliance rate per BU.\n"
            "- Shading Thresholds:\n"
            "   * Green: Non-PL Compliance >= 85.0%.\n"
            "   * Amber: Non-PL Compliance 70.0%-84.9%.\n"
            "   * Red: Non-PL Compliance < 70.0% (delayed regularization of sick/casual leaves).",

            "- MoM Delta: Change in timely post-facto leave submissions.\n"
            "- Driver Decomposition: Pinpoints BUs with delayed submission habits (leaves applied > 3 days after availing)."
        ),

        (
            "WFH Application Compliance\n(Category: Compliance | ID: wfh_application_compliance | Higher is Better)",
            "- Scope: Filters evaluated records where policy_event_type == 'WFH'.\n"
            "- Policy Rule: WFH requests must be submitted no later than 3 calendar days after availing (lead_time_days >= -3).\n"
            "- Formula: (Compliant WFH Events / Evaluable WFH Events) * 100%.\n"
            "- Trajectory: 3-month window trend direction and streak.",

            "- Evaluates WFH regularization compliance by Business Unit.\n"
            "- Gap: BU WFH Rate - Org WFH Benchmark (pp).\n"
            "- Ranking: Sorted descending by compliance.",

            "- Heatmap Grid: Monthly WFH compliance rates.\n"
            "- Color Thresholds:\n"
            "   * Green: WFH Compliance >= 85.0%.\n"
            "   * Amber: WFH Compliance 70.0%-84.9%.\n"
            "   * Red: WFH Compliance < 70.0%.",

            "- MoM Variance: Identifies whether remote work was pre-planned/timely regularized or logged as unapproved retroactive WFH."
        ),

        (
            "Attendance Exception Rate\n(Category: Attendance | ID: attendance_exception_rate | Lower is Better)",
            "- Source: Pre-aggregated employee-day facts table (1 row per employee per calendar date).\n"
            "- Eligible Days: Must have is_eligible_attendance_day == True (excludes scheduled holidays, weekly offs, and full-day approved leaves).\n"
            "- Evaluable Days: Eligible days excluding uncertain data quality records.\n"
            "- Numerator: Count of evaluable employee-days containing at least 1 attendance exception (unaccounted absence, missing swipe, tardiness, early departure, or under-hours).\n"
            "- Rate Formula: (Exception Employee-Days / Total Evaluable Eligible Employee-Days) * 100%.\n"
            "- Trajectory Direction (Lower is Better): Improving if slope <= -1.0 pp; Deteriorating if slope >= +1.0 pp; Stable otherwise.",

            "- Grouping: Aggregated per Business Unit for the latest month.\n"
            "- 3-Month Average: Mean exception rate across trailing 3 months.\n"
            "- Historical: Median pre-policy exception rate for that BU.\n"
            "- Gap vs Org: (BU Exception Rate) - (Org Benchmark Rate). Since lower is better, a negative gap is favorable (green) and a positive gap is adverse (red).\n"
            "- Ranking: Sorted ascending (lowest exception rate to highest).",

            "- Heatmap Grid: Monthly exception rate percentage per BU.\n"
            "- Color Coding (Lower is Better):\n"
            "   * Green (Healthy): Exception Rate < 10.0%.\n"
            "   * Amber (Moderate): Exception Rate 10.0%-25.0%.\n"
            "   * Red (Elevated): Exception Rate > 25.0%.",

            "- MoM Variance: Latest month exception rate minus previous month.\n"
            "- Driver Decomposition:\n"
            "   * BU Delta = Current BU Rate - Previous BU Rate.\n"
            "   * Weight = BU Eligible Days / Total Org Eligible Days.\n"
            "   * Highlights top BUs driving overall attendance discipline improvements or deterioration."
        ),

        (
            "Missing Swipe Rate\n(Category: Attendance | ID: missing_swipe_rate | Lower is Better)",
            "- Scope: Evaluated on eligible attendance days where employee is present/working.\n"
            "- Numerator: Employee-days where either first_in is missing or last_out is missing (single swipe / unrecorded exit).\n"
            "- Denominator: Total evaluable eligible employee-days.\n"
            "- Formula: (Missing Swipe Days / Evaluable Days) * 100%.\n"
            "- Trajectory: Lower is better (materiality threshold 1.0 pp).",

            "- Evaluates biometric swipe hygiene by Business Unit.\n"
            "- Gap: BU Missing Swipe Rate - Org Benchmark.\n"
            "- Ranking: Sorted ascending from cleanest swipe records to highest missing swipe rates.",

            "- Heatmap Matrix: Monthly missing swipe rates per BU.\n"
            "- Shading (Lower is Better):\n"
            "   * Green: Missing Swipe Rate < 5.0%.\n"
            "   * Amber: Missing Swipe Rate 5.0%-15.0%.\n"
            "   * Red: Missing Swipe Rate > 15.0%.",

            "- MoM Movement: Identifies teams improving clock-in/out discipline vs teams experiencing biometric reader compliance lapses."
        ),

        (
            "Absence Rate\n(Category: Attendance | ID: absence_rate | Lower is Better)",
            "- Scope: Eligible employee-days scheduled for work.\n"
            "- Numerator: Employee-days marked absent (both unaccounted absences and non-working absences).\n"
            "- Denominator: Total evaluable eligible employee-days.\n"
            "- Formula: (Absent Employee-Days / Evaluable Days) * 100%.\n"
            "- Trajectory: Lower is better. Threshold: 1.0 pp slope.",

            "- Evaluates overall absenteeism by Business Unit.\n"
            "- Gap: BU Absence Rate - Org Benchmark (pp).\n"
            "- Ranking: Sorted ascending (lowest absence rate to highest).",

            "- Heatmap Matrix: Monthly absence rate by BU.\n"
            "- Shading:\n"
            "   * Green: Absence Rate < 3.0%.\n"
            "   * Amber: Absence Rate 3.0%-8.0%.\n"
            "   * Red: Absence Rate > 8.0%.",

            "- MoM Variance: Details workforce shrinkage and unplanned absenteeism surges by BU."
        ),

        (
            "Attendance Regularisation Rate\n(Category: Attendance | ID: regularization_rate | Neutral)",
            "- Scope: Eligible employee-days with manual attendance regularization requests.\n"
            "- Formula: (Regularized Employee-Days / Evaluable Eligible Days) * 100%.\n"
            "- Direction: Neutral (monitored for operational reliance on manual overrides).",

            "- Displays BU regularization frequency vs company-wide average.\n"
            "- Gap: BU Rate - Org Rate.\n"
            "- Ranking: Sorted by deviation from normal range.",

            "- Heatmap Matrix: Monthly regularisation rates.\n"
            "- Color Coding: Blue divergence scale showing standard usage (< 8%) vs heavy reliance (> 15%).",

            "- Driver Decomposition: Pinpoints BUs experiencing sudden spikes in manual attendance corrections."
        ),

        (
            "Median Approval Turnaround\n(Category: Approval | ID: median_approval_turnaround | Lower is Better)",
            "- Scope: All approved leave & attendance requests grouped by application month.\n"
            "- Request Turnaround: (Approved On Date) - (Applied On Date) in calendar days (clipped to >= 0).\n"
            "- Metric Value: Median(turnaround_days) across all approved requests in the month.\n"
            "- Trajectory Direction (Lower is Better): Materiality threshold = 0.2 days. Improving if slope <= -0.2 d, Deteriorating if slope >= +0.2 d, else Stable.\n"
            "- Robustness: Median is resilient against severe outliers (e.g. neglected requests approved after 90 days).",

            "- Grouping: Aggregates approved requests by the reporting manager's Business Unit.\n"
            "- 3-Month Average: Mean of the BU's monthly median turnaround times.\n"
            "- Historical: Median pre-policy turnaround time for that BU.\n"
            "- Gap vs Org: BU Median Turnaround - Org Median Turnaround (in days). Negative gap indicates faster approval (favorable green badge).\n"
            "- Ranking: Sorted ascending (fastest turnaround to slowest).",

            "- Heatmap Grid: Monthly median approval turnaround (in days) per BU.\n"
            "- Shading (Lower is Better):\n"
            "   * Green (Rapid): Turnaround < 1.0 day.\n"
            "   * Amber (Normal): Turnaround 1.0-3.0 days.\n"
            "   * Red (Sluggish): Turnaround > 3.0 days.",

            "- MoM Net Movement: Change in median approval turnaround between consecutive months.\n"
            "- Driver Attribution: Decomposes overall management responsiveness changes by Business Unit."
        ),

        (
            "Average Approval Turnaround\n(Category: Approval | ID: average_approval_turnaround | Lower is Better)",
            "- Scope: All approved requests grouped by application month.\n"
            "- Value: Mean(turnaround_days) in calendar days.\n"
            "- Direction: Lower is better. Materiality threshold: 0.2 days.",

            "- Evaluates average management response time by Business Unit.\n"
            "- Gap: BU Average - Org Benchmark (days).\n"
            "- Ranking: Sorted ascending from fastest average approval to slowest.",

            "- Heatmap Grid: Monthly average turnaround per BU.\n"
            "- Shading:\n"
            "   * Green: Average < 1.5 days.\n"
            "   * Amber: Average 1.5-4.0 days.\n"
            "   * Red: Average > 4.0 days.",

            "- Driver Attribution: Identifies which BUs have long-tail approval bottlenecks pulling up the company mean."
        ),

        (
            "Pending Approval Count\n(Category: Approval | ID: pending_approval_count | Lower is Better)",
            "- Scope: All requests applied in the month with status Pending/Applied as of snapshot time.\n"
            "- Value: Integer count of unapproved requests.\n"
            "- Direction: Lower is better. Materiality threshold: 5 requests.",

            "- Current backlog count by Business Unit.\n"
            "- Gap: BU Pending Count - Org Benchmark.\n"
            "- Ranking: Sorted ascending from lowest backlog to highest backlog.",

            "- Heatmap Grid: Monthly pending counts per BU.\n"
            "- Shading:\n"
            "   * Green: 0-5 pending requests.\n"
            "   * Amber: 6-20 pending requests.\n"
            "   * Red: > 20 pending requests.",

            "- Driver Decomposition: Pinpoints BUs accumulating unapproved queues vs BUs actively clearing backlogs."
        ),

        (
            "Average Pending Approval Age\n(Category: Approval | ID: average_pending_approval_age | Lower is Better)",
            "- Scope: For pending requests: Age = (As-Of Date) - (Applied On Date) in days.\n"
            "- Value: Mean(pending_age_days).\n"
            "- Direction: Lower is better. Materiality threshold: 0.5 days.",

            "- Evaluates average backlog age per BU.\n"
            "- Gap: BU Pending Age - Org Benchmark (days).\n"
            "- Ranking: Sorted ascending (youngest queue to oldest queue).",

            "- Heatmap Grid: Monthly average pending age (days).\n"
            "- Shading:\n"
            "   * Green: Age < 2.0 days.\n"
            "   * Amber: Age 2.0-5.0 days.\n"
            "   * Red: Age > 5.0 days (stagnant approval queue).",

            "- Driver Attribution: Highlights managerial stagnation where pending items sit unresolved for weeks."
        ),

        (
            "Average Effective Hours\n(Category: Working Time | ID: average_effective_hours | Neutral)",
            "- Scope: Evaluated present days with valid recorded working duration in minutes.\n"
            "- Value: Mean(work_minutes) converted to 'Xh YYm' display format.\n"
            "- Direction: Neutral (monitors workload equilibrium and overtime patterns).",

            "- Average effective hours worked per present day by Business Unit.\n"
            "- Gap: BU Average - Org Average (in hours/minutes).\n"
            "- Ranking: Sorted by deviation from standard 8h 00m workday.",

            "- Heatmap Matrix: Monthly average hours per BU.\n"
            "- Shading: Balanced blue gradient (7h 45m-8h 45m normal; < 7h 30m under; > 9h 30m overtime risk).",

            "- Driver Attribution: Decomposes changes in working hours to identify over-extended or under-utilized teams."
        ),

        (
            "Median Effective Hours\n(Category: Working Time | ID: median_effective_hours | Neutral)",
            "- Value: Median(work_minutes) converted to 'Xh YYm'.\n"
            "- Robust against clock-out anomalies and extreme outliers.",

            "- Median workday length per Business Unit vs Org benchmark.\n"
            "- Ranking: Sorted by alignment to 8-hour workday standard.",

            "- Heatmap Matrix: Monthly median effective duration per BU.\n"
            "- Shading: Neutral blue distribution scale.",

            "- Driver Attribution: Tracks typical workday length movements excluding individual overtime extremes."
        ),

        (
            "Average Arrival Time\n(Category: Working Time | ID: average_arrival_time | Neutral)",
            "- Scope: Evaluated present days with valid in_time recorded.\n"
            "- Value: Mean(in_time_minutes_from_midnight) formatted into 12-hour AM/PM string (e.g., '09:42 AM').\n"
            "- Direction: Neutral.",

            "- Average morning clock-in time by Business Unit.\n"
            "- Gap: BU Arrival Time - Org Arrival Time (in minutes).\n"
            "- Ranking: Sorted from earliest arrival to latest arrival.",

            "- Heatmap Grid: Monthly average arrival times per BU.\n"
            "- Shading: Early (< 09:30 AM), Standard (09:30-10:15 AM), Late (> 10:15 AM).",

            "- Driver Attribution: Identifies which BUs shifted team arrival habits earlier or later."
        ),

        (
            "Median Arrival Time\n(Category: Working Time | ID: median_arrival_time | Neutral)",
            "- Value: Median(in_time_minutes_from_midnight) formatted into 'HH:MM AM/PM'.\n"
            "- Represents the typical arrival time for 50% of the team.",

            "- Median arrival time per BU vs company-wide median.\n"
            "- Ranking: Sorted by arrival order.",

            "- Heatmap Matrix: Monthly median arrival time distribution by BU.",

            "- Driver Attribution: Reflects structural core shift changes without noise from extreme late arrivals."
        ),

        (
            "Average Exit Time\n(Category: Working Time | ID: average_exit_time | Neutral)",
            "- Scope: Evaluated present days with valid out_time recorded.\n"
            "- Value: Mean(out_time_minutes_from_midnight) formatted into 12-hour AM/PM string (e.g., '06:48 PM').\n"
            "- Direction: Neutral.",

            "- Average departure time by Business Unit.\n"
            "- Gap: BU Exit Time - Org Exit Time (in minutes).\n"
            "- Ranking: Sorted from earliest departure to latest departure.",

            "- Heatmap Grid: Monthly average exit times per BU.\n"
            "- Shading: Early (< 06:00 PM), Standard (06:00-07:30 PM), Extended (> 07:30 PM).",

            "- Driver Attribution: Highlights teams frequently working late or leaving early."
        ),

        (
            "Median Exit Time\n(Category: Working Time | ID: median_exit_time | Neutral)",
            "- Value: Median(out_time_minutes_from_midnight) formatted into 'HH:MM AM/PM'.\n"
            "- Robust measure of typical end-of-day office departure.",

            "- Median departure time per BU vs company-wide median.\n"
            "- Ranking: Sorted by median departure time.",

            "- Heatmap Matrix: Monthly median departure distribution by BU.",

            "- Driver Attribution: Evaluates broader team shift movements vs individual late-night spikes."
        ),
    ]

    current_row = 3
    for idx, (m_name, traj_logic, ranking_logic, heatmap_logic, what_logic) in enumerate(rows_data):
        row_values = [m_name, traj_logic, ranking_logic, heatmap_logic, what_logic]
        ws.append(row_values)

        fill_color = c_zebra_even if idx % 2 == 0 else c_zebra_odd
        row_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")

        for col_idx in range(1, 6):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.fill = row_fill
            cell.border = border_thin
            cell.alignment = align_left

            if col_idx == 1:
                cell.font = font_metric
            else:
                cell.font = font_text

        ws.row_dimensions[current_row].height = 135
        current_row += 1

    col_widths = {
        1: 30,  # Matrix Name
        2: 52,  # Trajactory and Benchmark
        3: 44,  # Business unit benchmark ranking
        4: 44,  # Business Unit Trend Heatmap
        5: 46,  # What changes
    }
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"Generated: {out_path}")

if __name__ == "__main__":
    generate_workbook("output/Workforce_Trends_Metrics_Logic.xlsx")
    generate_workbook("Workforce_Trends_Metrics_Logic.xlsx")
