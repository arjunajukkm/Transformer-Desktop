"""
Tests for Attendance Regularization (A(R)) across Time and Leave Master
and Workforce Intelligence KPI Engine.

Validates that:
1. Records with Status 'A(R)' lacking Attendance Type are automatically populated
   with Attendance Type = 'Regularized', Approval Status = 'Approved', and admin audit defaults.
2. 'A(R)' records are excluded from the Absent Mailer sheet.
3. In the Dashboard KPI engine, 'A(R)' is never misclassified as Absent,
   and resolves to Present (Keka standard).
"""

import pytest
import pandas as pd
from datetime import date
from io import BytesIO
import openpyxl

from workforce_intelligence.normalization import normalize_attendance_category
from workforce_intelligence.kpi_engine import compute_workforce_intelligence_bundle
from time_leave_master import reconcile_time_and_leave


def test_normalize_attendance_category_ar():
    """Verify normalize_attendance_category correctly maps A(R) even if attendance type is blank or absent."""
    assert normalize_attendance_category("A(R)", "") == "Attendance Regularized"
    assert normalize_attendance_category("A(R)", None) == "Attendance Regularized"
    assert normalize_attendance_category("A(R)", "Absent") == "Attendance Regularized"
    assert normalize_attendance_category("P(R)", "") == "Attendance Regularized"
    assert normalize_attendance_category("AR", "") == "Attendance Regularized"
    assert normalize_attendance_category("A", "") == "Absent"
    assert normalize_attendance_category("A", "Absent") == "Absent"


def test_kpi_engine_ar_not_absent_and_is_present():
    """Verify user's exact case: FINBP009 with Status 'A(R)' resolves to Present, not Absent."""
    rows = [
        {
            "Employee Number": "FINBP009",
            "Employee Name": "Sanoop Sashindran",
            "Date": "2026-08-05",
            "Status": "A(R)",
            "Quantity": 1.0,
            "Approval Status": "Approved",
            "Attendance Type": "Regularized",
        }
    ]
    df = pd.DataFrame(rows)
    bundle = compute_workforce_intelligence_bundle(df)

    assert bundle["kpi_absent_days"] == 0.0
    assert bundle["kpi_3_present_days"] == 1.0
    assert bundle["kpi_9_attendance_exceptions_days"] == 1


def test_kpi_engine_ar_missing_attendance_type_not_absent():
    """Verify that when Attendance Type is completely missing, Status A(R) is not classified as Absent."""
    rows = [
        {
            "Employee Number": "FINBP009",
            "Employee Name": "Sanoop Sashindran",
            "Date": "2026-08-05",
            "Status": "A(R)",
            "Quantity": 1.0,
            "Approval Status": "Approved",
        }
    ]
    df = pd.DataFrame(rows)
    bundle = compute_workforce_intelligence_bundle(df)

    assert bundle["kpi_absent_days"] == 0.0
    assert bundle["kpi_3_present_days"] == 1.0
    assert bundle["kpi_9_attendance_exceptions_days"] == 1


def test_kpi_engine_ar_with_absent_attendance_type_not_absent():
    """Verify that even if Attendance Type was marked 'Absent' in raw dump, Status A(R) resolves to Present and 0 Absent."""
    rows = [
        {
            "Employee Number": "FINBP009",
            "Employee Name": "Sanoop Sashindran",
            "Date": "2026-08-05",
            "Status": "A(R)",
            "Attendance Type": "Absent",
            "Quantity": 1.0,
        }
    ]
    df = pd.DataFrame(rows)
    bundle = compute_workforce_intelligence_bundle(df)

    assert bundle["kpi_absent_days"] == 0.0
    assert bundle["kpi_3_present_days"] == 1.0
    assert bundle["kpi_9_attendance_exceptions_days"] == 1


def test_time_leave_master_reconciles_ar_and_excludes_from_absent_mailer(tmp_path):
    """Verify reconcile_time_and_leave assigns 'Regularized' to A(R) and omits it from Absent Mailer."""
    df_perf = pd.DataFrame([
        {
            "Employee Number": "FINBP009",
            "Employee Name": "Sanoop Sashindran",
            "Date": "2026-08-05",
            "Status": "A(R)",
            "Quantity": 1.0,
        },
        {
            "Employee Number": "FINBP010",
            "Employee Name": "John Doe",
            "Date": "2026-08-05",
            "Status": "A",
            "Quantity": 1.0,
        },
    ])

    perf_file = tmp_path / "perf.xlsx"
    df_perf.to_excel(perf_file, index=False)
    out_file = tmp_path / "output.xlsx"

    stats = reconcile_time_and_leave(
        perf_files=[str(perf_file)],
        output_path=str(out_file),
    )

    wb = openpyxl.load_workbook(str(out_file))
    assert "Daily Performance" in wb.sheetnames
    assert "Absent Mailer" in wb.sheetnames

    # Check Daily Performance sheet
    ws_perf = wb["Daily Performance"]
    headers = [cell.value for cell in ws_perf[1]]
    status_col_idx = headers.index("Status") + 1
    att_type_col_idx = headers.index("Attendance Type") + 1
    appr_status_col_idx = headers.index("Approval Status") + 1

    # Row 2 is FINBP009 (A(R))
    assert ws_perf.cell(row=2, column=status_col_idx).value == "A(R)"
    assert ws_perf.cell(row=2, column=att_type_col_idx).value == "Regularized"
    assert ws_perf.cell(row=2, column=appr_status_col_idx).value == "Approved"

    # Row 3 is FINBP010 (A)
    assert ws_perf.cell(row=3, column=status_col_idx).value == "A"

    # Check Absent Mailer sheet: FINBP009 (A(R)) MUST be excluded, only FINBP010 (A) should be present
    ws_mailer = wb["Absent Mailer"]
    mailer_rows = list(ws_mailer.iter_rows(values_only=True))
    assert len(mailer_rows) == 2  # 1 header + 1 data row
    mailer_headers = mailer_rows[0]
    emp_num_idx = mailer_headers.index("Employee Number")
    assert mailer_rows[1][emp_num_idx] == "FINBP010"


def test_keka_data_fetcher_reconciles_regularisations():
    """Verify KekaDataFetcher reconciles approved regularisation requests into A(R) and Regularized."""
    from keka_data_fetcher import KekaDataFetcher
    fetcher = KekaDataFetcher(subdomain="mock", client_id="mock", client_secret="mock", api_key="mock")

    emp_df = pd.DataFrame([
        {
            "Employee Number": "FINBP009",
            "Employee Name": "Sanoop Sashindran",
            "Job Title": "Analyst",
            "Business Unit": "Moshpit",
            "Department": "Finance",
            "Sub Department": "Billing",
            "Location": "Bangalore",
            "Reporting Manager": "Manager One",
        }
    ])

    # Raw attendance log from /time/attendance has no punches (normally yields 'A')
    mock_att_record = {
        "employeeNumber": "FINBP009",
        "attendanceDate": "2026-08-05T00:00:00Z",
        "dayType": 0,
        "totalEffectiveHours": 0.0,
        "totalGrossHours": 0.0,
        "totalBreakDuration": 0.0,
        "firstInOfTheDay": None,
        "lastOutOfTheDay": None,
    }

    # Approved regularisation request from /time/regularisationrequests
    mock_reg_record = {
        "employeeNumber": "FINBP009",
        "requestDate": "2026-08-05T00:00:00Z",
        "status": 2,  # Approved
        "requestedByName": "E Janani Sri",
        "requestedOn": "2026-09-18T09:57:06.163Z",
        "lastModified": "2026-09-18T09:57:06.163Z",
        "approvers": [],
    }

    # Mock _paginated_get
    def mock_paginated_get(endpoint, params=None, max_pages=None, label=None):
        if "/attendance" in endpoint:
            return [mock_att_record], None
        elif "/regularisationrequests" in endpoint:
            return [mock_reg_record], None
        return [], None

    fetcher._paginated_get = mock_paginated_get

    att_df, err = fetcher.fetch_attendance_records_api(
        from_date="2026-08-05",
        to_date="2026-08-05",
        emp_df=emp_df,
    )

    assert err is None
    assert len(att_df) == 1
    row = att_df.iloc[0]
    assert row["Employee Number"] == "FINBP009"
    assert row["Status"] == "A(R)"
    assert row["Attendance Type"] == "Regularized"
    assert row["Approval Status"] == "Approved"
    assert row["Applied By"] == "Admin"
    assert row["Approved By"] == "Admin"

